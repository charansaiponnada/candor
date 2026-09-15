# -*- coding: utf-8 -*-
"""Candor final pitch deck (LUMINIX'26 Grand Finale), OLED-black glass theme.

    python docs/pitch/build_deck.py   ->  docs/pitch/Candor_Pitch.pptx

Market figures are cited on-slide and on the Sources slide. Business-model
prices are proposals and labelled as such. The Product slide appears only when
screenshots exist in docs/screenshots/.
"""
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

ROOT = "C:/projects/candor"
TESTS = "133"  # flutter test, 15 Sep 2026

BLACK = C(0x00, 0x00, 0x00)
WHITE = C(0xFF, 0xFF, 0xFF)
ACCENT = C(0x00, 0xE5, 0xA0)
TIER2 = C(0x7C, 0x4D, 0xFF)
WARN = C(0xFF, 0xD9, 0x3D)
RED = C(0xFF, 0x6B, 0x6B)
DIM = C(0x9A, 0x9A, 0x9A)
FAINT = C(0x66, 0x66, 0x66)
LINE = C(0x2A, 0x2A, 0x2A)
GLASS = C(0x12, 0x12, 0x12)
GLASS2 = C(0x1B, 0x1B, 0x1B)
GREEN_FILL, GREEN_LINE = C(0x04, 0x16, 0x11), C(0x0E, 0x4A, 0x37)

FONT = "Segoe UI"
MONO = "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]
PAGE_NUMS = []  # filled in once the slide count is known


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLACK
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def text(s, x, y, w, h, runs, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         font=FONT, spacing=1.15, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, (t, o) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.line_spacing = o.get("spacing", spacing)
        if o.get("space_before"):
            p.space_before = Pt(o["space_before"])
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get("font", font)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return tb


def card(s, x, y, w, h, fill=GLASS, border=LINE, radius=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = border
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    sh.adjustments[0] = radius
    sh.text_frame.word_wrap = True
    return sh


def chip(s, x, y, w, h, label, color=ACCENT, size=11):
    sh = card(s, x, y, w, h, fill=GLASS2, border=color, radius=0.5)
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = color
    return sh


def eyebrow(s, label):
    text(s, 0.9, 0.62, 8, 0.3, label.upper(), size=12, color=ACCENT, bold=True)
    PAGE_NUMS.append(text(s, 11.5, 0.62, 0.95, 0.3, "-", size=11, color=DIM,
                          align=PP_ALIGN.RIGHT))
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(0.9), I(1.72), I(11.55), Pt(1))
    rule.fill.solid()
    rule.fill.fore_color.rgb = LINE
    rule.line.fill.background()
    rule.shadow.inherit = False


def title(s, t, sub=None):
    text(s, 0.9, 0.95, 11.5, 0.6, t, size=32, bold=True)
    if sub:
        text(s, 0.9, 1.92, 11.5, 0.4, sub, size=15, color=DIM)


def source(s, t):
    text(s, 0.9, 7.1, 11.55, 0.25, t, size=9, color=FAINT)


def arrow(s, x, y, w, h, color=ACCENT, shape=MSO_SHAPE.RIGHT_ARROW):
    a = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


# ---------------------------------------------------------------- TITLE
s = slide()
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, I(0.16), prs.slide_height)
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
bar.shadow.inherit = False

logo = ROOT + "/assets/icon/new_app_icon.png"
if os.path.exists(logo):
    s.shapes.add_picture(logo, I(0.85), I(1.3), height=I(1.6))

text(s, 0.9, 3.15, 11, 1.0, "CANDOR", size=76, bold=True, spacing=0.9)
text(s, 0.9, 4.25, 11, 0.5, "The assistant that tells you when it's holding back.",
     size=24, color=ACCENT)
text(s, 0.9, 4.95, 9.5, 0.9,
     "A privacy-first AI assistant that runs two quantized small language models entirely "
     "on your phone. No server, no subscription, no network permission.",
     size=15, color=DIM, spacing=1.3)

for i, (lab, col) in enumerate([("100% on-device", ACCENT), ("0 network calls", ACCENT),
                                ("%s/%s tests pass" % (TESTS, TESTS), TIER2),
                                ("Runs on a mid-range phone", WARN)]):
    chip(s, 0.9 + i * 2.75, 6.1, 2.5, 0.42, lab, col)

text(s, 0.9, 6.95, 11.5, 0.3,
     "LUMINIX'26 Grand Finale   \u00b7   Charan Sai Ponnada   \u00b7   github.com/charansaiponnada/candor",
     size=11, color=FAINT, font=MONO)

# ------------------------------------------------------------- PROBLEM
s = slide()
eyebrow(s, "The problem")
title(s, "Cloud AI disappears exactly when you need it.",
      "Villages, tunnels, flights, outages, an empty data pack \u2014 and the assistant is gone.")

card(s, 0.9, 2.35, 5.55, 4.1, fill=C(0x14, 0x0B, 0x0B), border=C(0x3A, 0x1E, 0x1E))
text(s, 1.25, 2.62, 4.9, 0.3, "TODAY \u2014 RENTED INTELLIGENCE", size=12, bold=True, color=RED)
text(s, 1.25, 3.15, 4.9, 3.1, [
    ("No network, no assistant.", {"bold": True, "size": 15}),
    ("Every answer is a round-trip to someone else's datacentre.", {"size": 13, "color": DIM}),
    ("Your words leave the device.", {"bold": True, "size": 15, "space_before": 12}),
    ("Prompts travel, get logged, get retained.", {"size": 13, "color": DIM}),
    ("Premium phones first.", {"bold": True, "size": 15, "space_before": 12}),
    ("Built-in on-device AI ships mostly on expensive flagships.", {"size": 13, "color": DIM}),
    ("Metered and rate-limited.", {"bold": True, "size": 15, "space_before": 12}),
    ("Stop paying and the intelligence switches off.", {"size": 13, "color": DIM}),
], spacing=1.15)

arrow(s, 6.62, 4.15, 0.55, 0.5, ACCENT)

card(s, 7.0, 2.35, 5.45, 4.1, fill=GREEN_FILL, border=GREEN_LINE)
text(s, 7.35, 2.62, 4.8, 0.3, "CANDOR \u2014 OWNED INTELLIGENCE", size=12, bold=True, color=ACCENT)
text(s, 7.35, 3.15, 4.8, 3.1, [
    ("Two models ride in the APK.", {"bold": True, "size": 15}),
    ("Qwen2.5 0.5B + 1.5B, 4-bit quantized, run by llama.cpp.", {"size": 13, "color": DIM}),
    ("Incapable of phoning home.", {"bold": True, "size": 15, "space_before": 12}),
    ("The release manifest declares no INTERNET permission.", {"size": 13, "color": DIM}),
    ("Runs on the phone you own.", {"bold": True, "size": 15, "space_before": 12}),
    ("Proven on a mid-range Samsung Galaxy, CPU only.", {"size": 13, "color": DIM}),
    ("Honest about its own limits.", {"bold": True, "size": 15, "space_before": 12}),
    ("Every reply names the model that answered, and why.", {"size": 13, "color": DIM}),
], spacing=1.15)

text(s, 0.9, 6.75, 11.55, 0.35,
     "Airplane mode is the demo: chat, calculators, converters, notes, timers and drafts all still work.",
     size=13, color=ACCENT, align=PP_ALIGN.CENTER)

# -------------------------------------------------------------- MARKET
s = slide()
eyebrow(s, "Why now")
title(s, "A billion Indians online \u2014 not everywhere.",
      "The people who most need a helpful assistant are the ones cloud AI serves worst.")

stats = [("48 vs 127", "internet subscriptions per 100 people, rural vs urban India (Q1 2026)", WARN, 30),
         ("300M", "people worldwide still live outside any mobile broadband coverage", RED, 38),
         ("84%", "of GenAI users worry the data they enter could go public", TIER2, 38),
         ("13%", "of phones shipped in India in H1 2025 had built-in GenAI", ACCENT, 38)]
for i, (big, cap, col, sz) in enumerate(stats):
    x = 0.9 + i * 2.95
    card(s, x, 2.35, 2.7, 1.9)
    text(s, x + 0.22, 2.55, 2.35, 0.75, big, size=sz, bold=True, color=col)
    text(s, x + 0.22, 3.4, 2.35, 0.8, cap, size=11, color=DIM, spacing=1.2)

card(s, 0.9, 4.55, 5.6, 2.35)
text(s, 1.2, 4.78, 5.0, 0.3, "THE MARKET", size=12, bold=True, color=ACCENT)
text(s, 1.2, 5.18, 5.0, 0.6, "$11.8B \u2192 $56.8B", size=30, bold=True, color=WHITE)
text(s, 1.2, 5.9, 5.0, 0.9, [
    ("Global edge-AI market, 2025 \u2192 2030 (36.9% CAGR).", {}),
    ("Smartphones are 80.5% of edge-AI hardware by volume.", {"space_before": 4}),
], size=12.5, color=DIM, spacing=1.2)

card(s, 6.85, 4.55, 5.6, 2.35, fill=GREEN_FILL, border=GREEN_LINE)
text(s, 7.15, 4.78, 5.0, 0.3, "OUR INSIGHT", size=12, bold=True, color=ACCENT)
text(s, 7.15, 5.18, 5.0, 1.6, [
    ("87% of new phones in India ship without built-in GenAI.", {"bold": True, "size": 16,
                                                                 "color": WHITE}),
    ("Candor doesn't wait for a flagship or a signal: a 4-bit model cascade gives any "
     "mid-range Android a private, offline assistant today.", {"size": 12.5, "color": DIM,
                                                                "space_before": 8}),
], spacing=1.2)

source(s, "Sources: TRAI via The Week (2026) \u00b7 GSMA SOMIC 2025 \u00b7 Cisco Privacy Survey "
          "2024 \u00b7 Counterpoint (H1 2025) \u00b7 BCC Research (2025) \u00b7 MarketsandMarkets "
          "\u2014 full list on the last slide")

# ---------------------------------------------------------------- USERS
s = slide()
eyebrow(s, "Who it's for")
title(s, "Built for where the network doesn't reach.",
      "Each segment has the same two needs: it must work offline, and the data must stay put.")

segments = [
    ("Frontline field workers", "Health, agri-extension and survey staff in low-signal villages: "
     "notes, unit maths and quick reference with zero uploads.", ACCENT),
    ("Students off the grid", "A patient explainer that works without a data pack, on the family's "
     "shared mid-range phone.", ACCENT),
    ("Travel, defence & disaster", "Flights, remote postings and outages \u2014 exactly when "
     "networks are down or not permitted.", WARN),
    ("Privacy-bound professionals", "Doctors, lawyers, journalists: drafts and questions that "
     "must never touch a third-party server.", TIER2),
    ("Locked-down enterprises", "Field sales and ops fleets where IT policy forbids sending data "
     "to cloud AI.", TIER2),
    ("Everyday users", "Anyone tired of subscriptions, rate limits and wondering where their "
     "prompts end up.", WARN),
]
for i, (h, b, col) in enumerate(segments):
    x = 0.9 + (i % 3) * 3.93
    y = 2.4 + (i // 3) * 2.2
    card(s, x, y, 3.6, 1.95)
    text(s, x + 0.28, y + 0.25, 3.1, 0.35, h, size=15, bold=True, color=col)
    text(s, x + 0.28, y + 0.72, 3.1, 1.1, b, size=12, color=DIM, spacing=1.2)

# ----------------------------------------------------------- ARCHITECTURE
s = slide()
eyebrow(s, "How it works")
title(s, "One honest router, two tiers, zero packets.",
      "Every query is answered at the cheapest tier that can hold it \u2014 and labelled with which one did.")


def node(x, y, w, h, head, body, col=WHITE, fill=GLASS, bsize=11):
    card(s, x, y, w, h, fill=fill, border=col)
    text(s, x + 0.12, y + 0.16, w - 0.24, 0.3, head, size=13, bold=True, color=col,
         align=PP_ALIGN.CENTER)
    lines = [(t, {"align": PP_ALIGN.CENTER}) for t in body.split("\n")]
    text(s, x + 0.12, y + 0.5, w - 0.24, h - 0.58, lines, size=bsize, color=DIM,
         align=PP_ALIGN.CENTER, spacing=1.1)


node(0.9, 3.5, 1.85, 1.15, "USER QUERY", "typed or spoken\n(on-device STT)", ACCENT)
arrow(s, 2.85, 3.94, 0.35, 0.28)
node(3.3, 3.35, 2.0, 1.45, "INTENT MATCH", "regex classifier\nskill / action / model", WARN)

node(5.85, 2.45, 2.5, 1.0, "SKILLS ENGINE", "calculator \u00b7 units \u00b7 dates \u00b7 notes \u00b7 dice",
     WARN, bsize=10)
node(5.85, 3.60, 2.5, 1.0, "PLATFORM ACTIONS", "timer \u00b7 apps \u00b7 SMS \u00b7 email \u00b7 settings",
     WARN, bsize=10)
node(5.85, 4.75, 2.5, 1.1, "ROUTER", "multi-step question?\nTier-1 unsure / refusing?", ACCENT)
arrow(s, 5.45, 2.82, 0.3, 0.26, WARN)
arrow(s, 5.45, 3.97, 0.3, 0.26, WARN)
arrow(s, 5.45, 5.17, 0.3, 0.26, ACCENT)

node(8.9, 4.35, 1.9, 1.9, "TIER-1 \u00b7 0.5B",
     "fast default\nself-scored answer\n\nskipped for\nmulti-step asks", ACCENT)
arrow(s, 8.45, 5.17, 0.35, 0.26, ACCENT)

node(11.3, 4.35, 1.6, 0.85, "TIER-2 \u00b7 1.5B", "deeper answer", TIER2, bsize=10)
node(11.3, 5.4, 1.6, 0.85, "CONSTRAINED", "Tier-1 + reason", RED, bsize=10)
arrow(s, 10.9, 4.65, 0.3, 0.26, TIER2)
arrow(s, 10.9, 5.7, 0.3, 0.26, RED)

card(s, 8.6, 6.55, 2.5, 0.62, fill=C(0x10, 0x10, 0x18), border=TIER2)
text(s, 8.68, 6.64, 2.34, 0.5, [
    ("DEVICE GATE", {"size": 10.5, "bold": True, "color": TIER2, "align": PP_ALIGN.CENTER}),
    ("battery \u2265 20%  \u00b7  thermal < MODERATE",
     {"size": 9, "color": DIM, "align": PP_ALIGN.CENTER}),
])
arrow(s, 9.72, 6.28, 0.26, 0.25, TIER2, MSO_SHAPE.UP_ARROW)

card(s, 0.9, 5.2, 4.4, 1.75, fill=GLASS2, border=LINE)
text(s, 1.2, 5.4, 3.8, 0.3, "THE OUTPUT CONTRACT", size=11, bold=True, color=ACCENT)
text(s, 1.2, 5.78, 3.9, 0.35, "\u25cf Tier-2 \u00b7 1.5B answering \u00b7 12s", size=14, bold=True,
     color=WHITE, font=MONO)
text(s, 1.2, 6.2, 3.9, 0.7,
     "The chat shows which model is working, live. Comparisons, step-by-step and code go "
     "straight to Tier-2 \u2014 no wasted 0.5B pass.",
     size=10.5, color=DIM, spacing=1.15)

# ------------------------------------------------------------- PRODUCT
shots = [(f, cap) for f, cap in [
    ("chat_tier1.png", "Tier-1 answer with its honesty pill"),
    ("chat_tier2_streaming.png", "Escalated to Tier-2, streaming live"),
    ("tier_sheet.png", "Tap the pill: why this model answered"),
    ("skills.png", "Deterministic skills, zero tokens"),
] if os.path.exists("%s/docs/screenshots/%s" % (ROOT, f))]
if shots:
    s = slide()
    eyebrow(s, "The product")
    title(s, "Real screens, captured on the phone.")
    h = 4.5  # sized by height: phone shots are 1080x2340 and must clear the captions
    w = h * 1080 / 2340
    gap = (11.55 - w * len(shots)) / max(len(shots) - 1, 1)
    for i, (f, cap) in enumerate(shots):
        x = 0.9 + i * (w + gap) if len(shots) > 1 else 5.5
        s.shapes.add_picture("%s/docs/screenshots/%s" % (ROOT, f), I(x), I(2.2), height=I(h))
        text(s, x - 0.45, 6.85, w + 0.9, 0.5, cap, size=11, color=DIM, align=PP_ALIGN.CENTER)

# --------------------------------------------------------------- PROOF
s = slide()
eyebrow(s, "Evidence")
title(s, "Proven, not promised.",
      "Everything below is measurable on a mid-range handset in airplane mode, today.")

stats = [("%s/%s" % (TESTS, TESTS), "automated tests pass\nrouter \u00b7 skills \u00b7 stores \u00b7 UI", ACCENT),
         ("0", "INTERNET permissions\nin the release manifest", ACCENT),
         ("~1.5 GB", "two 4-bit models in the APK\n0.5B + 1.5B, Q4_K_M", TIER2),
         ("100%", "offline capability\nno feature needs a network", WARN)]
for i, (big, cap, col) in enumerate(stats):
    x = 0.9 + i * 2.95
    card(s, x, 2.35, 2.7, 1.75)
    text(s, x + 0.22, 2.55, 2.3, 0.7, big, size=36, bold=True, color=col)
    text(s, x + 0.22, 3.35, 2.3, 0.7, cap, size=11, color=DIM, spacing=1.2)

text(s, 0.9, 4.5, 5.6, 0.3, "IN THE BOX", size=12, bold=True, color=ACCENT)
text(s, 0.9, 4.95, 5.6, 2.2, [
    ("Models & Benchmarks \u2014 live load time, tokens/sec and size per bundled model, "
     "measured on this phone.", {}),
    ("Prompt Lab \u2014 per-tier temperature, top-k/top-p and penalty sliders with a live sandbox.",
     {"space_before": 9}),
    ("Debug panel \u2014 real battery and thermal readout, plus a simulated constrained device.",
     {"space_before": 9}),
    ("Voice input, chat history, per-tier model pickers, Vulkan GPU toggle.",
     {"space_before": 9}),
], size=12.5, color=C(0xCC, 0xCC, 0xCC), spacing=1.2)

text(s, 7.0, 4.5, 5.45, 0.3, "THE STACK", size=12, bold=True, color=ACCENT)
rows = [("Flutter / Dart", "single codebase, custom glass widget set"),
        ("llama.cpp", "via the llama_flutter_android plugin"),
        ("Qwen2.5 Q4_K_M", "0.5B Tier-1 + 1.5B Tier-2, in the APK"),
        ("Kotlin channel", "battery, thermal, model-file prep"),
        ("Vulkan offload", "optional GPU acceleration")]
for i, (k, v) in enumerate(rows):
    y = 4.95 + i * 0.44
    text(s, 7.0, y, 2.4, 0.3, k, size=12, bold=True, color=WHITE)
    text(s, 9.5, y, 2.95, 0.3, v, size=11.5, color=DIM)

# ---------------------------------------------------------- COMPETITION
s = slide()
eyebrow(s, "Competitive landscape")
title(s, "Offline, private, affordable — pick all three.")

YES, PART, NO = ("\u25cf", ACCENT), ("\u25d0", WARN), ("\u25cb", FAINT)
cols = ["Works fully\noffline", "Data stays\non device", "Runs on mid-\nrange phones",
        "Says which model\nanswered", "Exact skills &\nphone actions"]
rows = [("Cloud assistants", "ChatGPT, Gemini app", [NO, NO, YES, NO, PART]),
        ("Built-in phone AI", "Galaxy AI, Gemini Nano", [PART, PART, NO, NO, YES]),
        ("Local LLM apps", "AI Edge Gallery, PocketPal", [YES, YES, PART, NO, PART]),
        ("Candor", "this project", [YES, YES, YES, YES, YES])]
x0, cw, y0, rh = 4.3, 1.63, 2.3, 0.98
for j, c in enumerate(cols):
    text(s, x0 + j * cw, y0, cw, 0.6, c, size=11.5, bold=True, color=DIM, align=PP_ALIGN.CENTER)
for i, (name, eg, marks) in enumerate(rows):
    y = y0 + 0.75 + i * rh
    mine = name == "Candor"
    card(s, 0.9, y, 11.55, rh - 0.14, fill=GREEN_FILL if mine else GLASS,
         border=GREEN_LINE if mine else LINE)
    text(s, 1.2, y + 0.14, 3.0, 0.35, name, size=15, bold=True, color=ACCENT if mine else WHITE)
    text(s, 1.2, y + 0.48, 3.0, 0.3, eg, size=10.5, color=DIM)
    for j, (sym, col) in enumerate(marks):
        text(s, x0 + j * cw, y + 0.12, cw, 0.6, sym, size=26, color=col, align=PP_ALIGN.CENTER)

source(s, "\u25cf yes   \u25d0 partial / device- or feature-dependent   \u25cb no.   "
          "Based on public product documentation, Sep 2026.")

# ------------------------------------------------------- BUSINESS MODEL
s = slide()
eyebrow(s, "Business model")
title(s, "Three ways to earn \u2014 and almost no cost to serve.",
      "Proposed model. Inference runs on the user's own phone, so there is no per-query GPU bill.")

plans = [
    ("CONSUMERS", "Free + one-time Pro", ACCENT,
     ["Free: the full offline assistant, both tiers, every skill.",
      "Pro (one-time, ~\u20b9149 proposed): extra model packs, document Q&A, custom skills.",
      "No subscription \u2014 nothing to meter."]),
    ("ORGANISATIONS", "Offline-AI licence", TIER2,
     ["NGOs, health & agri programs, field fleets, defence.",
      "Per-device annual licence: private deployment, custom skill packs, MDM-ready builds.",
      "Priced per seat, not per query."]),
    ("OEMs & TELCOS", "Preload royalty", WARN,
     ["Budget-phone makers ship an AI assistant with no cloud contract.",
      "Small per-unit royalty; differentiates sub-premium devices.",
      "Distribution at the scale of shipments."]),
]
for i, (tag, head, col, bullets) in enumerate(plans):
    x = 0.9 + i * 3.93
    card(s, x, 2.4, 3.6, 3.35)
    text(s, x + 0.3, 2.62, 3.0, 0.3, tag, size=11, bold=True, color=col)
    text(s, x + 0.3, 2.95, 3.0, 0.45, head, size=19, bold=True, color=WHITE)
    text(s, x + 0.3, 3.55, 3.05, 2.1,
         [("\u2022  " + b, {"space_before": 6 if k else 0}) for k, b in enumerate(bullets)],
         size=12, color=DIM, spacing=1.2)

card(s, 0.9, 6.0, 11.55, 0.95, fill=GREEN_FILL, border=GREEN_LINE)
text(s, 1.3, 6.2, 10.8, 0.6, [
    ("Unit economics: a cloud assistant pays for every token it generates. Candor's marginal "
     "cost per answer is \u2248 \u20b90 \u2014 the phone already paid for the compute.",
     {"size": 14, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}),
], spacing=1.2)

# --------------------------------------------------------- GO-TO-MARKET
s = slide()
eyebrow(s, "Go-to-market")
title(s, "Start where the network ends.",
      "Plan: land with mission-driven field programs, prove it, then scale through devices.")

phases = [("0\u20133 MONTHS", "Launch", ACCENT,
           ["Open beta: GitHub release + Play Store listing.",
            "Opt-in, on-device escalation log to tune the router.",
            "First Indic-language model tier."]),
          ("3\u20139 MONTHS", "Pilot", TIER2,
           ["2\u20133 pilots with NGOs or health / agri field programs.",
            "Custom skill packs per program.",
            "Measure answers delivered offline."]),
          ("9\u201318 MONTHS", "Scale", WARN,
           ["Organisation licences from pilot results.",
            "OEM / telco preload conversations.",
            "Document Q&A (on-device RAG) as Pro."])]
for i, (when, head, col, bullets) in enumerate(phases):
    x = 0.9 + i * 3.93
    card(s, x, 2.4, 3.6, 3.1)
    text(s, x + 0.3, 2.62, 3.0, 0.3, when, size=11, bold=True, color=col)
    text(s, x + 0.3, 2.95, 3.0, 0.45, head, size=22, bold=True, color=WHITE)
    text(s, x + 0.3, 3.6, 3.05, 1.8,
         [("\u2022  " + b, {"space_before": 6 if k else 0}) for k, b in enumerate(bullets)],
         size=12.5, color=DIM, spacing=1.2)
    if i < 2:
        arrow(s, x + 3.64, 3.8, 0.25, 0.3, col)

text(s, 0.9, 5.8, 11.55, 0.3, "WHAT WE WILL MEASURE", size=12, bold=True, color=ACCENT)
for i, (k, v) in enumerate([("Offline answers / week", "the north-star"),
                            ("Escalation rate", "how often Tier-2 is needed"),
                            ("Battery per answer", "cost to the user's phone"),
                            ("Pilot retention", "do field teams keep using it")]):
    x = 0.9 + i * 2.95
    text(s, x, 6.2, 2.8, 0.3, k, size=13, bold=True, color=WHITE)
    text(s, x, 6.52, 2.8, 0.3, v, size=11, color=DIM)

# ------------------------------------------------------------- ROADMAP
s = slide()
eyebrow(s, "What's next")
title(s, "Same privacy guarantee, more capability.",
      "The hard part \u2014 offline inference behind an honest router \u2014 already ships.")

items = [("Indic languages", "A multilingual tier so rural users can ask in their own language.", ACCENT),
         ("On-device RAG", "Index notes and documents locally. Grounded answers, zero uploads.", TIER2),
         ("Learned router", "Train the escalation decision on real on-device logs, not keywords.", WARN),
         ("Bigger tiers", "3B+ quantized models as phones catch up. Same cascade, smarter answers.", ACCENT),
         ("More skills", "Calendar, files and offline reference packs. Deterministic, unhallucinatable.", TIER2),
         ("Play Store", "Ship the release pipeline that already builds the APK today.", WARN)]
for i, (h, b, col) in enumerate(items):
    x = 0.9 + (i % 3) * 3.93
    y = 2.4 + (i // 3) * 1.75
    card(s, x, y, 3.6, 1.5)
    text(s, x + 0.28, y + 0.22, 3.1, 0.3, h, size=15, bold=True, color=col)
    text(s, x + 0.28, y + 0.64, 3.1, 0.8, b, size=12, color=DIM, spacing=1.15)

text(s, 0.9, 6.15, 11.55, 0.6,
     "Known limits we state openly: English-only today \u00b7 confidence is self-reported "
     "(the plugin exposes no logprobs) \u00b7 one model in RAM at a time.",
     size=12, color=DIM, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- TEAM
s = slide()
eyebrow(s, "The team")
title(s, "One builder, end to end.",
      "Design, Flutter app, on-device inference, router, tests and this deck — built solo.")

photo = ROOT + "/docs/pitch/charan.png"
ring = s.shapes.add_shape(MSO_SHAPE.OVAL, I(0.85), I(2.35), I(2.7), I(2.7))
ring.fill.background()
ring.line.color.rgb = ACCENT
ring.line.width = Pt(2)
ring.shadow.inherit = False
if os.path.exists(photo):
    s.shapes.add_picture(photo, I(0.95), I(2.45), width=I(2.5))
text(s, 0.9, 5.25, 3.2, 0.4, "Charan Sai Ponnada", size=20, bold=True)
text(s, 0.9, 5.7, 3.2, 0.6, "Software Engineer (AI), Aynstyn Technologies · B.Tech AI & DS, VRSEC",
     size=11.5, color=DIM, spacing=1.2)
text(s, 0.9, 6.45, 3.3, 0.6, [
    ("github.com/charansaiponnada", {}),
    ("linkedin.com/in/charansaiponnada", {"space_before": 3}),
], size=10.5, color=ACCENT, font=MONO)

creds = [
    ("Edge AI, shipped before", "Primary author, IEEE ISAECT 2025: a fine-tuned vision-language "
     "model deployed on a Raspberry Pi to guide visually impaired pedestrians.", ACCENT),
    ("Honest LLMs is the research", "Sole-author paper under review at IEEE InCODE 2026: semantic "
     "consistency as an unsupervised hallucination signal.", TIER2),
    ("Production LLM reliability", "At Aynstyn: black-box test suites that made LLM responses "
     "reliable across 3+ production workflows.", WARN),
    ("Proven under pressure", "2nd place among the Top 10 finalists at YUVAAN 2026, IIT Hyderabad "
     "— out of 7,600+ registrants.", ACCENT),
]
for i, (h, b, col) in enumerate(creds):
    x = 4.3 + (i % 2) * 4.2
    y = 2.4 + (i // 2) * 2.1
    card(s, x, y, 3.95, 1.85)
    text(s, x + 0.28, y + 0.25, 3.4, 0.35, h, size=15, bold=True, color=col)
    text(s, x + 0.28, y + 0.7, 3.4, 1.05, b, size=12, color=DIM, spacing=1.2)

text(s, 4.3, 6.75, 8.15, 0.35,
     "Also building a 3D-chromatin-aware genomic foundation model (Hopfield-Mamba) from scratch.",
     size=11, color=FAINT)

# --------------------------------------------------------------- CLOSE
s = slide()
eyebrow(s, "The ask")
title(s, "Private AI on the phone people already own.")

asks = [("Pilot partners", "Field programs in low-connectivity regions willing to trial Candor.", ACCENT),
        ("Mentorship & internship", "Help taking Candor from hackathon build to Play Store.", TIER2),
        ("Your hardest question", "Try it in airplane mode \u2014 we'll show which model answers.", WARN)]
for i, (h, b, col) in enumerate(asks):
    x = 0.9 + i * 3.93
    card(s, x, 2.4, 3.6, 1.9)
    text(s, x + 0.3, 2.65, 3.0, 0.35, h, size=17, bold=True, color=col)
    text(s, x + 0.3, 3.15, 3.0, 1.0, b, size=13, color=DIM, spacing=1.2)

card(s, 0.9, 4.75, 11.55, 2.1, fill=GREEN_FILL, border=GREEN_LINE)
text(s, 1.3, 5.1, 10.8, 1.6, [
    ("\u201cAn assistant that knows when to hold back \u2014 and tells you.\u201d",
     {"size": 26, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}),
    ("Thank you.", {"size": 18, "color": ACCENT, "align": PP_ALIGN.CENTER, "space_before": 10}),
    ("github.com/charansaiponnada/candor",
     {"size": 12, "color": DIM, "align": PP_ALIGN.CENTER, "space_before": 8, "font": MONO}),
])

# ------------------------------------------------------------- SOURCES
s = slide()
eyebrow(s, "Appendix \u00b7 Sources")
title(s, "Where the numbers come from.")
refs = [
    ("48 vs 127 internet subscriptions / 100 (rural vs urban), 1.09B users, Q1 2026",
     "TRAI data, reported by The Week, 23 Jun 2026 \u2014 theweek.in/news/biz-tech/2026/06/23/india-telecom-growth-digital-divide.html"),
    ("~300 million people (4% of the world) outside mobile broadband coverage",
     "GSMA, The State of Mobile Internet Connectivity 2025, as quoted by connectivity.technology (Oct 2025)"),
    ("84% of GenAI users concerned their data could go public",
     "Cisco 2024 Consumer Privacy Survey \u2014 newsroom.cisco.com (Oct 2024)"),
    ("13% GenAI smartphones in India H1 2025 (3% H1 2024)",
     "Counterpoint Research, reported by Business Standard, Sep 2025"),
    ("Edge AI $11.8B (2025) \u2192 $56.8B (2030), 36.9% CAGR",
     "BCC Research press release via GlobeNewswire, 2 Oct 2025"),
    ("Smartphones 80.5% of edge-AI hardware volume (2024)",
     "MarketsandMarkets, Edge AI Hardware Market report"),
]
for i, (claim, src) in enumerate(refs):
    y = 2.3 + i * 0.75
    text(s, 0.9, y, 11.55, 0.3, claim, size=13, bold=True, color=WHITE)
    text(s, 0.9, y + 0.32, 11.55, 0.3, src, size=10.5, color=DIM, font=MONO)

# ------------------------------------------------------------- numbering
total = len(prs.slides)
for i, tb in enumerate(PAGE_NUMS):
    tb.text_frame.paragraphs[0].runs[0].text = "%d / %d" % (i + 2, total)

out = ROOT + "/docs/pitch/Candor_Pitch.pptx"
prs.save(out)
print("saved", out, total, "slides", os.path.getsize(out), "bytes")
